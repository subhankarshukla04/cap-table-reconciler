"""Excel export smoke tests — verify the workbook is valid + contains the
expected key strings."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.corpus import Corpus
from src.engine import resolve
from src.export import verdict_to_xlsx
from src.models import Query

FIXTURES = Path(__file__).parent / "fixtures"
RULES = Path(__file__).parent.parent / "src" / "rules"


@pytest.fixture(scope="module")
def live_corpus() -> Corpus:
    return Corpus.load(RULES)


def _verdict_for(name: str, corpus):
    raw = json.loads((FIXTURES / f"{name}.json").read_text())
    raw = {k: v for k, v in raw.items() if not k.startswith("_")}
    return resolve(Query.model_validate(raw), corpus)


def test_xlsx_export_loads_back(live_corpus):
    v = _verdict_for("praxis", live_corpus)
    payload = verdict_to_xlsx(v)
    wb = load_workbook(io.BytesIO(payload), read_only=False)
    assert "Verdict" in wb.sheetnames
    assert "Citations" in wb.sheetnames


def test_xlsx_verdict_sheet_has_query_inputs(live_corpus):
    v = _verdict_for("praxis", live_corpus)
    wb = load_workbook(io.BytesIO(verdict_to_xlsx(v)))
    ws = wb["Verdict"]
    body = "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert "SG" in body and "IN" in body
    assert "exercise" in body
    assert "perquisite" in body.lower()


def test_xlsx_citations_sheet_lists_all_sources(live_corpus):
    v = _verdict_for("solstice", live_corpus)
    wb = load_workbook(io.BytesIO(verdict_to_xlsx(v)))
    ws = wb["Citations"]
    body = "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert "IRAS" in body
    assert "IR21" in body
    assert "Income Tax Act 1947" in body


def test_xlsx_all_three_fixtures_export_cleanly(live_corpus):
    for name in ("praxis", "pelaut", "solstice"):
        v = _verdict_for(name, live_corpus)
        payload = verdict_to_xlsx(v)
        # Must be a valid xlsx (zip header)
        assert payload[:2] == b"PK"
