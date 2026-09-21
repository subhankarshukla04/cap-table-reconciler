"""Batch CSV mode.

Parse a CSV of employee scenarios, resolve each against the loaded corpus,
and emit a combined audit-grade Excel workbook with:

  - 'Summary' sheet: one row per employee with primary verdict, secondary
    perspectives, computed taxable amount, weakest confidence.
  - 'Citations' sheet: full citation chain for every rule fired across all rows.
  - 'Inputs' sheet: echoed input data for audit trail.

CSV schema (header row required):
    employee_name, parent_jurisdiction, employee_jurisdiction, event,
    tax_status, grant_date, vesting_date, event_date, fmv_at_event,
    exercise_price, options_in_event

employee_name is optional (used for output labelling only — engine ignores it).
All date fields ISO-8601 (YYYY-MM-DD). Numeric fields blank-allowed.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date
from typing import IO, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .corpus import Corpus
from .engine import resolve
from .models import Event, Jurisdiction, Query, TaxStatus, Verdict

EXPECTED_COLUMNS = [
    "employee_name",
    "parent_jurisdiction",
    "employee_jurisdiction",
    "event",
    "tax_status",
    "grant_date",
    "vesting_date",
    "event_date",
    "fmv_at_event",
    "exercise_price",
    "options_in_event",
]


@dataclass
class BatchRow:
    employee_name: str
    query: Query
    verdict: Verdict
    error: str | None = None


class BatchError(Exception):
    """Raised when the CSV is malformed."""


def _parse_date(s: str | None) -> date | None:
    if not s or not s.strip():
        return None
    try:
        return date.fromisoformat(s.strip())
    except ValueError as e:
        raise BatchError(f"Invalid date '{s}' — expected YYYY-MM-DD") from e


def _parse_float(s: str | None) -> float | None:
    if s is None or not s.strip():
        return None
    try:
        return float(s.strip())
    except ValueError as e:
        raise BatchError(f"Invalid number '{s}'") from e


def _parse_int(s: str | None) -> int | None:
    if s is None or not s.strip():
        return None
    try:
        return int(s.strip())
    except ValueError as e:
        raise BatchError(f"Invalid integer '{s}'") from e


def parse_csv(fp: IO[str]) -> list[tuple[str, Query | None, str | None]]:
    """Parse the CSV into (employee_name, Query or None, error_message or None).

    A row that fails validation is returned with error_message set — the batch
    runner still produces a row in the output workbook so the user can fix it.
    """
    reader = csv.DictReader(fp)
    if reader.fieldnames is None:
        raise BatchError("CSV is empty — expected a header row")
    missing = [c for c in EXPECTED_COLUMNS if c not in reader.fieldnames]
    # employee_name is optional; allow without it
    missing = [c for c in missing if c != "employee_name"]
    if missing:
        raise BatchError(
            f"CSV missing required columns: {', '.join(missing)}. "
            f"Expected: {', '.join(EXPECTED_COLUMNS)}"
        )

    out: list[tuple[str, Query | None, str | None]] = []
    for i, row in enumerate(reader, start=2):  # row 2 = first data line
        name = (row.get("employee_name") or f"Row {i - 1}").strip() or f"Row {i - 1}"
        try:
            q = Query(
                parent_jurisdiction=Jurisdiction(row["parent_jurisdiction"].strip()),
                employee_jurisdiction=Jurisdiction(row["employee_jurisdiction"].strip()),
                event=Event(row["event"].strip()),
                tax_status=TaxStatus(row.get("tax_status", "resident").strip() or "resident"),
                grant_date=_parse_date(row.get("grant_date")),
                vesting_date=_parse_date(row.get("vesting_date")),
                event_date=_parse_date(row.get("event_date")),
                fmv_at_event=_parse_float(row.get("fmv_at_event")),
                exercise_price=_parse_float(row.get("exercise_price")),
                options_in_event=_parse_int(row.get("options_in_event")),
            )
            out.append((name, q, None))
        except BatchError as e:
            out.append((name, None, str(e)))
        except (ValueError, Exception) as e:  # pydantic ValidationError etc.
            out.append((name, None, str(e)))
    return out


def run_batch(parsed: Iterable[tuple[str, Query | None, str | None]], corpus: Corpus) -> list[BatchRow]:
    out: list[BatchRow] = []
    for name, q, err in parsed:
        if q is None:
            out.append(BatchRow(employee_name=name, query=None, verdict=None, error=err))  # type: ignore[arg-type]
            continue
        verdict = resolve(q, corpus)
        out.append(BatchRow(employee_name=name, query=q, verdict=verdict))
    return out


# ---- Excel rendering -------------------------------------------------------


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
ERROR_FILL = PatternFill("solid", fgColor="F5C6CB")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="45556A")
BODY_FONT = Font(name="Calibri", size=10, color="16202C")
MONO_FONT = Font(name="Menlo", size=10, color="16202C")
THIN = Side(border_style="thin", color="D8DEE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SUMMARY_HEADERS = [
    "Employee",
    "Parent",
    "Employee jur.",
    "Event",
    "Status",
    "Event date",
    "Options",
    "FMV",
    "Strike",
    "Computed taxable",
    "Primary taxing",
    "Primary treatment",
    "Withholding",
    "Secondary perspectives",
    "Weakest confidence",
    "Error",
]
SUMMARY_WIDTHS = [22, 8, 10, 12, 12, 12, 10, 10, 10, 18, 12, 18, 13, 28, 16, 30]


def batch_to_xlsx(rows: list[BatchRow]) -> bytes:
    wb = Workbook()
    _sheet_summary(wb.active, rows)
    _sheet_citations(wb.create_sheet("Citations"), rows)
    _sheet_inputs(wb.create_sheet("Inputs"), rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_summary(ws, rows: list[BatchRow]) -> None:
    ws.title = "Summary"
    ws["A1"] = "ESOP Atlas — Batch Verdict Summary"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    ws.merge_cells("A1:P1")
    ws["A2"] = f"Generated {date.today().isoformat()} — {len(rows)} row(s)"
    ws["A2"].font = Font(name="Calibri", size=9, color="45556A", italic=True)
    ws.merge_cells("A2:P2")

    header_row = 4
    for col, (h, w) in enumerate(zip(SUMMARY_HEADERS, SUMMARY_WIDTHS), start=1):
        cell = ws.cell(row=header_row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[header_row].height = 32

    for i, row in enumerate(rows, start=1):
        r = header_row + i
        if row.error or row.query is None:
            ws.cell(row=r, column=1, value=row.employee_name).font = BODY_FONT
            for col in range(2, len(SUMMARY_HEADERS)):
                ws.cell(row=r, column=col, value="—").font = BODY_FONT
            err_cell = ws.cell(row=r, column=len(SUMMARY_HEADERS), value=row.error or "Unknown error")
            err_cell.font = BODY_FONT
            err_cell.fill = ERROR_FILL
            err_cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in range(1, len(SUMMARY_HEADERS) + 1):
                ws.cell(row=r, column=col).border = BORDER
            ws.row_dimensions[r].height = 26
            continue

        q = row.query
        v = row.verdict
        values = [
            row.employee_name,
            q.parent_jurisdiction.value,
            q.employee_jurisdiction.value,
            q.event.value.replace("_", " "),
            q.tax_status.value.replace("_", " "),
            q.event_date.isoformat() if q.event_date else "—",
            f"{q.options_in_event:,}" if q.options_in_event is not None else "—",
            f"{q.fmv_at_event:,.2f}" if q.fmv_at_event is not None else "—",
            f"{q.exercise_price:,.2f}" if q.exercise_price is not None else "—",
            f"{v.computed_amount_taxable:,.2f}" if v.computed_amount_taxable is not None else "—",
            v.primary_rule.taxing_jurisdiction.value,
            v.primary_rule.tax_treatment.value.replace("_", " "),
            v.primary_rule.withholding,
            "; ".join(
                f"{r.taxing_jurisdiction.value}: {r.tax_treatment.value.replace('_', ' ')}"
                for r in v.secondary_rules
            ) or "—",
            v.weakest_confidence.value.replace("_", " "),
            "",
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = MONO_FONT if col in (7, 8, 9, 10) else BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
        ws.row_dimensions[r].height = 26


def _sheet_citations(ws, rows: list[BatchRow]) -> None:
    ws.title = "Citations"
    headers = ["Employee", "Rule ID", "Taxing jur.", "Authority", "Reference", "Source URL", "Retrieved"]
    widths = [22, 32, 12, 22, 60, 65, 12]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w

    r = 2
    for row in rows:
        if not row.verdict:
            continue
        for rule in [row.verdict.primary_rule, *row.verdict.secondary_rules]:
            for c in rule.citations:
                ws.cell(row=r, column=1, value=row.employee_name).font = BODY_FONT
                ws.cell(row=r, column=2, value=rule.rule_id).font = MONO_FONT
                ws.cell(row=r, column=3, value=rule.taxing_jurisdiction.value).font = BODY_FONT
                ws.cell(row=r, column=4, value=c.authority).font = BODY_FONT
                ws.cell(row=r, column=5, value=c.reference).font = BODY_FONT
                ws.cell(row=r, column=6, value=c.source_url or "—").font = BODY_FONT
                ws.cell(row=r, column=7, value=c.retrieved_on.isoformat() if c.retrieved_on else "—").font = MONO_FONT
                for col in range(1, 8):
                    ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
                    ws.cell(row=r, column=col).border = BORDER
                ws.row_dimensions[r].height = 32
                r += 1


def _sheet_inputs(ws, rows: list[BatchRow]) -> None:
    ws.title = "Inputs"
    for col, (h, w) in enumerate(zip(EXPECTED_COLUMNS, [22, 10, 12, 12, 12, 12, 12, 12, 12, 14, 14]), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w

    for i, row in enumerate(rows, start=2):
        q = row.query
        ws.cell(row=i, column=1, value=row.employee_name).font = BODY_FONT
        if q is None:
            continue
        values = [
            q.parent_jurisdiction.value,
            q.employee_jurisdiction.value,
            q.event.value,
            q.tax_status.value,
            q.grant_date.isoformat() if q.grant_date else "",
            q.vesting_date.isoformat() if q.vesting_date else "",
            q.event_date.isoformat() if q.event_date else "",
            q.fmv_at_event if q.fmv_at_event is not None else "",
            q.exercise_price if q.exercise_price is not None else "",
            q.options_in_event if q.options_in_event is not None else "",
        ]
        for col, val in enumerate(values, start=2):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = MONO_FONT if col >= 9 else BODY_FONT
            cell.border = BORDER


def sample_csv() -> str:
    """Return a sample CSV the user can download as a template."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPECTED_COLUMNS)
    w.writerow(["Priya Sharma", "SG", "IN", "exercise", "resident", "2024-01-01", "2025-01-01", "2026-04-01", "0.50", "0.10", "4000"])
    w.writerow(["Wei Lin", "IN", "SG", "exercise", "resident", "2024-03-01", "2025-03-01", "2026-04-15", "4.00", "2.00", "20000"])
    w.writerow(["Atlas Engineer", "US", "IN", "exercise", "resident", "2023-09-01", "2024-09-01", "2026-04-30", "4.00", "0.50", "8000"])
    w.writerow(["Dubai Founder", "AE", "IN", "exercise", "resident", "2025-01-01", "2026-01-01", "2026-05-01", "6.00", "1.00", "10000"])
    w.writerow(["HK Engineer", "HK", "IN", "exercise", "resident", "2024-04-01", "2025-04-01", "2026-04-30", "3.00", "0.30", "6000"])
    return buf.getvalue()
