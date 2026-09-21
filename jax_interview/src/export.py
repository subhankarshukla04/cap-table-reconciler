"""Verdict exporters — Excel (openpyxl) and print-ready HTML.

PDF is intentionally NOT generated natively (weasyprint adds heavy native
dependencies; libpango etc.). The print-ready HTML view uses CSS @page rules
so the user can Cmd+P → Save as PDF from the browser. Demo-grade and
zero-dependency.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Rule, Verdict


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="45556A")
BODY_FONT = Font(name="Calibri", size=10, color="16202C")
MONO_FONT = Font(name="Menlo", size=10, color="16202C")
THIN = Side(border_style="thin", color="D8DEE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def verdict_to_xlsx(verdict: Verdict) -> bytes:
    """Render a Verdict to a polished, audit-grade Excel workbook.

    Two sheets:
      - 'Verdict': query inputs + summary + rule-by-rule treatment
      - 'Citations': flat list of every citation across primary + secondary rules
    """
    wb = Workbook()
    _sheet_verdict(wb.active, verdict)
    _sheet_citations(wb.create_sheet("Citations"), verdict)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_verdict(ws, verdict: Verdict) -> None:
    ws.title = "Verdict"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    # Title row
    ws["A1"] = "ESOP Atlas — Compliance Verdict"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    ws.merge_cells("A1:B1")

    ws["A2"] = f"Generated {date.today().isoformat()}"
    ws["A2"].font = Font(name="Calibri", size=9, color="45556A", italic=True)
    ws.merge_cells("A2:B2")

    row = 4
    q = verdict.query
    pairs = [
        ("Parent jurisdiction", q.parent_jurisdiction.value),
        ("Employee jurisdiction", q.employee_jurisdiction.value),
        ("Event", q.event.value.replace("_", " ")),
        ("Tax status", q.tax_status.value.replace("_", " ")),
        ("Grant date", q.grant_date.isoformat() if q.grant_date else "—"),
        ("Vesting date", q.vesting_date.isoformat() if q.vesting_date else "—"),
        ("Event date", q.event_date.isoformat() if q.event_date else "—"),
        ("FMV at event", f"{q.fmv_at_event:,.2f}" if q.fmv_at_event is not None else "—"),
        ("Exercise price", f"{q.exercise_price:,.2f}" if q.exercise_price is not None else "—"),
        ("Options in event", f"{q.options_in_event:,}" if q.options_in_event is not None else "—"),
        ("Computed amount taxable", f"{verdict.computed_amount_taxable:,.2f}" if verdict.computed_amount_taxable is not None else "—"),
        ("Weakest confidence", verdict.weakest_confidence.value.replace("_", " ")),
    ]
    _write_table(ws, row, "Query inputs", pairs)
    row += len(pairs) + 3

    # Primary rule
    row = _write_rule_block(ws, row, verdict.primary_rule, "Primary rule — employee-jurisdiction tax view")
    for sec in verdict.secondary_rules:
        row = _write_rule_block(ws, row, sec, f"Secondary rule — {sec.taxing_jurisdiction.value} perspective")

    # Footer
    ws.cell(row=row + 1, column=1, value="Note").font = LABEL_FONT
    ws.cell(row=row + 1, column=2, value="Engine output is deterministic and rule-citable. See 'Citations' tab for the full source list. Audit-grade compliance verdict, not tax advice.").font = BODY_FONT
    ws.cell(row=row + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row + 1].height = 36


def _write_table(ws, start_row: int, title: str, pairs: list[tuple[str, str]]) -> None:
    ws.cell(row=start_row, column=1, value=title).font = HEADER_FONT
    ws.cell(row=start_row, column=1).fill = HEADER_FILL
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)
    for i, (k, v) in enumerate(pairs, start=1):
        r = start_row + i
        ws.cell(row=r, column=1, value=k).font = LABEL_FONT
        ws.cell(row=r, column=2, value=v).font = (
            MONO_FONT if any(c.isdigit() for c in v) and "," in v else BODY_FONT
        )
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=2).border = BORDER


def _write_rule_block(ws, start_row: int, rule: Rule, heading: str) -> int:
    ws.cell(row=start_row, column=1, value=heading).font = HEADER_FONT
    ws.cell(row=start_row, column=1).fill = HEADER_FILL
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=2)

    pairs = [
        ("Rule ID", rule.rule_id),
        ("Tax treatment", rule.tax_treatment.value.replace("_", " ")),
        ("Rate", rule.rate_description),
        ("Withholding", f"{rule.withholding}" + (f" — {rule.withholding_note}" if rule.withholding_note else "")),
        ("Filing", rule.filing_requirement or "—"),
        ("Documents", "; ".join(rule.documents_required) if rule.documents_required else "—"),
        ("Confidence", rule.confidence.value.replace("_", " ")),
    ]
    for i, (k, v) in enumerate(pairs, start=1):
        r = start_row + i
        ws.cell(row=r, column=1, value=k).font = LABEL_FONT
        ws.cell(row=r, column=2, value=v).font = BODY_FONT
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=2).border = BORDER
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 22 if len(v) < 80 else 44)

    caveats_row = start_row + len(pairs) + 1
    if rule.caveats:
        ws.cell(row=caveats_row, column=1, value="Caveats").font = LABEL_FONT
        ws.cell(row=caveats_row, column=2, value="\n".join(f"• {c}" for c in rule.caveats)).font = BODY_FONT
        ws.cell(row=caveats_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[caveats_row].height = max(30, 16 * len(rule.caveats))
        return caveats_row + 2
    return caveats_row + 1


def _sheet_citations(ws, verdict: Verdict) -> None:
    ws.title = "Citations"
    headers = ["Rule ID", "Authority", "Reference", "Source URL", "Retrieved", "Note"]
    widths = [28, 22, 60, 70, 12, 50]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w

    row = 2
    for r in [verdict.primary_rule, *verdict.secondary_rules]:
        for c in r.citations:
            ws.cell(row=row, column=1, value=r.rule_id).font = MONO_FONT
            ws.cell(row=row, column=2, value=c.authority).font = BODY_FONT
            ws.cell(row=row, column=3, value=c.reference).font = BODY_FONT
            ws.cell(row=row, column=4, value=c.source_url or "—").font = BODY_FONT
            ws.cell(row=row, column=5, value=c.retrieved_on.isoformat() if c.retrieved_on else "—").font = MONO_FONT
            ws.cell(row=row, column=6, value=c.note or "—").font = BODY_FONT
            for col in range(1, 7):
                ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(row=row, column=col).border = BORDER
            ws.row_dimensions[row].height = 38
            row += 1
