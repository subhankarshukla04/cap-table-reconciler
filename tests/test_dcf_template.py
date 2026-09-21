"""DCF template tests (W4.3)."""

from __future__ import annotations

import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.dcf_template import build_dcf_template, build_dcf_template_bytes
from src.models import (
    AntiDilution,
    AntiDilutionVariant,
    CapTable,
    Company,
    LiquidationPreference,
    LPType,
    ShareClass,
    ShareClassType,
)


def _ct():
    return CapTable(
        company=Company(name="DCF Co", currency="USD"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common, shares_outstanding=5_000_000),
            ShareClass(
                name="Series A",
                type=ShareClassType.preferred,
                shares_outstanding=1_000_000,
                issue_price=1.0, issue_date=date(2024, 1, 1), seniority_rank=1,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=1_000_000, type=LPType.non_participating),
                anti_dilution=AntiDilution(variant=AntiDilutionVariant.broad_based_weighted_average),
            ),
        ],
    )


def test_template_has_required_sheets_in_order():
    wb = build_dcf_template(_ct())
    expected_order = ["Read Me", "Assumptions", "Projection",
                      "Terminal Value", "Equity Bridge", "Per-Class FV"]
    assert wb.sheetnames == expected_order


def test_template_defines_named_ranges_for_assumptions():
    wb = build_dcf_template(_ct())
    names = {n.name for n in wb.defined_names.values()}
    for required in ("base_revenue", "wacc", "terminal_g", "tax_rate",
                     "ebitda_margin", "sum_pv_fcff", "pv_terminal_value",
                     "enterprise_value", "equity_value_total"):
        assert required in names, f"missing named range: {required}"


def test_template_per_class_fv_references_sidecar_slugs():
    """W5.4: references must use Excel external-link syntax so they
    resolve across the sidecar workbook."""
    wb = build_dcf_template(_ct())
    ws = wb["Per-Class FV"]
    common_share_cell = ws.cell(row=2, column=3).value
    # Old form `=share_count_Common` only resolved locally, leaving
    # #NAME? on first open. The fixed form names the sidecar file
    # so Excel can resolve cross-workbook.
    assert common_share_cell == "='[dcf_sidecar.xlsx]Cap Inputs'!share_count_Common"


def test_template_external_ref_filename_is_customisable():
    """The sidecar filename can be overridden."""
    wb = build_dcf_template(_ct(), sidecar_filename="my_sidecar.xlsx")
    ws = wb["Per-Class FV"]
    common_share_cell = ws.cell(row=2, column=3).value
    assert "my_sidecar.xlsx" in common_share_cell


def test_template_bytes_round_trip_through_openpyxl():
    raw = build_dcf_template_bytes(_ct())
    wb = openpyxl.load_workbook(BytesIO(raw))
    assert "Per-Class FV" in wb.sheetnames
    assert "Projection" in wb.sheetnames
    # Verify the projection formulas are stored (cell.value starts with =).
    ws = wb["Projection"]
    rev_y1 = ws.cell(row=2, column=2).value
    assert isinstance(rev_y1, str) and rev_y1.startswith("=")


def test_template_yellow_input_cells_have_default_values():
    """The template ships sensible defaults so the analyst can see the
    DCF chain compute end-to-end before they override."""
    wb = build_dcf_template(_ct())
    ws = wb["Assumptions"]
    # Base revenue (row 2 col 2) is yellow with a default value
    assert ws.cell(row=2, column=2).value == 10_000_000
    # WACC default
    for r in range(2, 20):
        if ws.cell(row=r, column=3).value == "wacc":
            assert ws.cell(row=r, column=2).value == 0.12
            return
    pytest.fail("WACC slug not found in Assumptions sheet")
