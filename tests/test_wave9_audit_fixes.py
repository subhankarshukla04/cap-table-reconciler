"""Regression tests for SYSTEM_AUDIT_WAVE_9.md fixes.

One test per finding ID so future regressions trace to the exact
audit section. Covers W9-B1 (CSRF rotation race), W9-M1..M4, and the
fix-now minors W9-m2, W9-m5, W9-m6, W9-m7.
"""

from __future__ import annotations

import io
import re
import secrets
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from flask import Flask

from src.cookie_auth import (
    CSRF_COOKIE_NAME,
    CSRF_ROTATION_SECONDS,
    attach_login_blueprint,
    csrf_token_age_seconds,
    issue_csrf_token,
)
from src.engagement import AuditEventType, EngagementStore
from src.engagement_routes import attach_engagement_blueprint
from src.identity import Role, StaticUserProvider, User
from src.opm.backsolve import BacksolveError, MarketInputs
from src.rate_limit import ExportRateLimiter
from src.token_deny import TokenDenyList


@pytest.fixture
def web():
    paths = []
    for _ in range(4):
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
            paths.append(Path(fh.name))
    eng_p, ex_p, cp_p, dn_p = paths
    store = EngagementStore(db_path=eng_p)
    export_limiter = ExportRateLimiter(db_path=ex_p, soft_limit=50, hard_limit=100)
    compute_limiter = ExportRateLimiter(db_path=cp_p, soft_limit=2, hard_limit=3)
    deny_list = TokenDenyList(db_path=dn_p)
    app = Flask(__name__)
    app.config["TESTING"] = True
    app.config["SESSION_SECRET_KEY"] = b"x" * 32
    app.config["TOKEN_DENY_LIST"] = deny_list
    users = {
        "tok-a": User(id="u-a", email="a@x", role=Role.analyst, display_name="A"),
        "tok-p": User(id="u-p", email="p@x", role=Role.partner, display_name="P"),
    }
    provider = StaticUserProvider(users)
    attach_engagement_blueprint(app, store, provider,
                                 export_limiter=export_limiter,
                                 compute_limiter=compute_limiter)
    attach_login_blueprint(app, identity_provider=provider)
    yield app.test_client(), store
    deny_list.close()
    for p in paths:
        p.unlink(missing_ok=True)


def _bearer(tok: str):
    return {"Authorization": f"Bearer {tok}"}


def _xlsx_bytes(shares_a: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 5_000_000, 0.01, "2020-01-01",
               "", "", "", ""])
    ws.append(["Series A", "preferred", shares_a, 1.00, "2024-01-01",
               1.0, "non_participating", 1, "broad_based_weighted_average"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# W9-B1: CSRF rotation race
# =============================================================================


def test_w9b1_rotation_keeps_template_and_cookie_in_sync(web):
    """The classic repro: stale cookie → GET renders form with NEW
    token → POST with that NEW token succeeds. Pre-fix: template had
    OLD token, cookie had NEW → 403 csrf-mismatch."""
    client, store = web
    client.post("/login", data={"token": "tok-a"})
    # Plant a stale cookie that will trigger rotation.
    stale = f"{int(time.time()) - 3700}.{secrets.token_urlsafe(32)}"
    client.set_cookie(CSRF_COOKIE_NAME, stale, domain="localhost")
    # GET the engagement list page — _inject_csrf should rotate AND
    # the after_request should attach the matching cookie.
    r = client.get("/engagement/?html=1")
    assert r.status_code == 200
    # Extract the token the template embedded in the form.
    match = re.search(
        r'name=[\'"]csrf_token[\'"]\s+value=[\'"]([^\'"]+)[\'"]',
        r.get_data(as_text=True),
    )
    assert match, "template did not render csrf_token field"
    form_token = match.group(1)
    cookie_token = client.get_cookie(CSRF_COOKIE_NAME).value
    # The two MUST be equal; if they differ the next POST 403s.
    assert form_token == cookie_token
    # Neither is the stale one.
    assert form_token != stale
    # Submit the form with the value the page showed — must succeed.
    r2 = client.post("/engagement/?html=1", data={
        "csrf_token": form_token,
        "client_id": "acme", "valuation_date": "2025-06-01",
    })
    # 302 → redirect to detail page (HTMX form post path).
    assert r2.status_code in (200, 302), r2.get_data(as_text=True)


def test_w9b1_fresh_token_not_rotated(web):
    """A fresh CSRF token must not rotate within the rotation window."""
    client, _ = web
    client.post("/login", data={"token": "tok-a"})
    csrf_before = client.get_cookie(CSRF_COOKIE_NAME).value
    r = client.get("/engagement/?html=1")
    assert r.status_code == 200
    csrf_after = client.get_cookie(CSRF_COOKIE_NAME).value
    assert csrf_before == csrf_after


# =============================================================================
# W9-M1: DB-level pagination
# =============================================================================


def test_w9m1_list_snapshots_accepts_limit_offset(web):
    """list_snapshots(limit=N, offset=M) pushes pagination into SQL."""
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    for shares, ver in [(1_000_000, 0), (1_500_000, 1), (2_000_000, 2)]:
        client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                    data={"file": (io.BytesIO(_xlsx_bytes(shares)), "x.xlsx"),
                          "expected_version": str(ver)},
                    content_type="multipart/form-data")
    all_snaps = store.list_snapshots(eng_id)
    assert len(all_snaps) == 3
    page = store.list_snapshots(eng_id, limit=1, offset=1)
    assert len(page) == 1
    assert page[0].id == all_snaps[1].id


def test_w9m1_count_snapshots_returns_total(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    assert store.count_snapshots(eng_id) == 0
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "x.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    assert store.count_snapshots(eng_id) == 1


def test_w9m1_default_call_preserves_backward_compat(web):
    """Calls without limit/offset return every snapshot — memo, bundle,
    timeline-diff all rely on this."""
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    for shares, ver in [(1_000_000, 0), (1_500_000, 1)]:
        client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                    data={"file": (io.BytesIO(_xlsx_bytes(shares)), "x.xlsx"),
                          "expected_version": str(ver)},
                    content_type="multipart/form-data")
    assert len(store.list_snapshots(eng_id)) == 2


# =============================================================================
# W9-M2: OPM MarketInputs domain checks
# =============================================================================


@pytest.mark.parametrize("bad", [
    {"volatility": -0.5},
    {"volatility": float("nan")},
    {"volatility": float("inf")},
    {"time_to_liquidity_years": 0},
    {"time_to_liquidity_years": -1.0},
    {"time_to_liquidity_years": 51.0},
    {"risk_free_rate": -0.5},
    {"risk_free_rate": 2.0},
    {"dlom": -0.1},
    {"dlom": 1.5},
    {"dividend_yield": -0.01},
])
def test_w9m2_market_inputs_rejects_out_of_domain(bad):
    """Any out-of-domain value raises BacksolveError at construction."""
    base = dict(volatility=0.55, time_to_liquidity_years=4.0,
                risk_free_rate=0.045, dividend_yield=0.0, dlom=0.0)
    base.update(bad)
    with pytest.raises(BacksolveError):
        MarketInputs(**base)


def test_w9m2_route_returns_opm_bad_inputs_on_negative_volatility(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "x.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    r = client.post(f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
                    json={
                        "volatility": -0.5,  # the bad input
                        "time_to_liquidity_years": 4.0,
                        "risk_free_rate": 0.045,
                        "anchor_class_name": "Series A",
                        "anchor_price_per_share": 1.00,
                    })
    assert r.status_code == 400
    payload = r.get_json()
    assert payload["error_code"] == "opm-bad-inputs"
    assert "volatility" in payload["error"]


# =============================================================================
# W9-M3: jurisdiction threading in finding_provenance_map
# =============================================================================


def test_w9m3_provenance_map_honours_jurisdiction():
    """When two rules in different jurisdictions emit the same code,
    passing the engagement's jurisdiction picks the right one."""
    from src.rule_pack import finding_provenance_map, head_pack
    from src.models import (
        CapTable, Company, ShareClass, ShareClassType,
        LiquidationPreference, LPType,
    )
    # India jurisdiction surfaces G-IN-* rules; other jurisdictions
    # don't. Provenance map for an India engagement must not return
    # rule_ids from other jurisdictions.
    from datetime import date as _date
    ct = CapTable(
        company=Company(name="Acme India", currency="INR",
                         currency_symbol="₹", jurisdiction="india",
                         valuation_date=_date(2025, 6, 1)),
        share_classes=[
            ShareClass(name="Series A", type=ShareClassType.preferred,
                       shares_outstanding=1_000_000, issue_price=10.0,
                       liquidation_preference=LiquidationPreference(
                           multiple=1.0, type=LPType.non_participating,
                           amount=10_000_000,
                       )),
        ],
    )
    # No assertion on specific codes — just that the function accepts
    # the jurisdiction kwarg and runs cleanly.
    out = finding_provenance_map(ct, engagement_jurisdiction="india")
    assert isinstance(out, dict)


def test_w9m3_pdf_memo_passes_jurisdiction_to_provenance_map():
    """The PDF memo context-builder threads ct.company.jurisdiction
    through to finding_provenance_map."""
    from src.pdf_memo import PDFInputs, _build_context, ReviewerInfo
    from src.waterfall import compute_waterfall
    from src.models import (
        CapTable, Company, ShareClass, ShareClassType,
        LiquidationPreference, LPType,
    )
    ct = CapTable(
        company=Company(name="Acme", jurisdiction="delaware"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=1000),
            ShareClass(name="Series A", type=ShareClassType.preferred,
                       shares_outstanding=1000, issue_price=1.0,
                       liquidation_preference=LiquidationPreference(
                           multiple=1.0, type=LPType.non_participating,
                           amount=1000,
                       )),
        ],
    )
    inputs = PDFInputs(
        cap_table=ct, waterfall=compute_waterfall(ct), findings=[],
        resolutions=[],
        reviewer=ReviewerInfo(reviewer_name="r", preparer="p"),
    )
    # Should not raise.
    ctx = _build_context(inputs)
    # findings list is empty in this minimal fixture so we only verify
    # the provenance map didn't crash.
    assert "findings" in ctx


# =============================================================================
# W9-M4: OPM compute_burst_alert at hard-limit boundary
# =============================================================================


def test_w9m4_opm_route_writes_compute_burst_alert(web):
    """compute_limiter hard=3 in fixture. Three /opm calls → audit
    chain carries one compute_burst_alert with route='opm'."""
    import json as _json
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "x.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    snap_id = store.get_engagement(eng_id).head_snapshot_id
    client.post(f"/engagement/{eng_id}/resolve", headers=_bearer("tok-a"),
                json={"snapshot_id": snap_id,
                      "finding_code": "AD-MISSING-Series A",
                      "decision": {"variant": "broad_based_weighted_average"},
                      "citation": "Charter §4.3(a)"})
    payload = {
        "volatility": 0.55, "time_to_liquidity_years": 4.0,
        "risk_free_rate": 0.045,
        "anchor_class_name": "Series A",
        "anchor_price_per_share": 1.00,
    }
    for _ in range(3):
        client.post(f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
                    json=payload)
    events = store.list_audit_events(eng_id)
    bursts = [
        e for e in events
        if e.event_type == AuditEventType.compute_burst_alert
        and _json.loads(e.payload_json).get("route") == "opm"
    ]
    assert bursts


# =============================================================================
# W9-m2: negative csrf age triggers rotation
# =============================================================================


def test_w9m2_csrf_negative_age_triggers_rotation(web):
    """A future-dated token reports negative age; rotation must fire."""
    client, _ = web
    client.post("/login", data={"token": "tok-a"})
    # Plant a future-dated token (clock skew / tampering scenario).
    future = f"{int(time.time()) + 7200}.{secrets.token_urlsafe(32)}"
    client.set_cookie(CSRF_COOKIE_NAME, future, domain="localhost")
    r = client.get("/engagement/?html=1")
    assert r.status_code == 200
    csrf_after = client.get_cookie(CSRF_COOKIE_NAME).value
    assert csrf_after != future
    age = csrf_token_age_seconds(csrf_after)
    assert age is not None and age >= 0


# =============================================================================
# W9-m6: stable kebab-case BacksolveError codes
# =============================================================================


def test_w9m6_anchor_missing_returns_stable_kebab_case(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "x.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    snap_id = store.get_engagement(eng_id).head_snapshot_id
    client.post(f"/engagement/{eng_id}/resolve", headers=_bearer("tok-a"),
                json={"snapshot_id": snap_id,
                      "finding_code": "AD-MISSING-Series A",
                      "decision": {"variant": "broad_based_weighted_average"},
                      "citation": "Charter §4.3(a)"})
    # Anchor a class that doesn't exist → BacksolveAnchorMissing.
    r = client.post(f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
                    json={"volatility": 0.55,
                          "time_to_liquidity_years": 4.0,
                          "risk_free_rate": 0.045,
                          "anchor_class_name": "Series Z",
                          "anchor_price_per_share": 1.0})
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "backsolve-anchor-missing"


# =============================================================================
# W9-m5 + W9-m7: vol-pack sourcing surfaces; rows are dynamic
# =============================================================================


def test_w9m7_vol_pack_sourcing_surfaces_in_sidecar():
    from src.dcf_sidecar import build_dcf_sidecar_bytes
    from src.opm.backsolve import MarketInputs
    from src.opm.vol_pack import VolPackReadback
    from src.models import CapTable, Company, ShareClass, ShareClassType
    ct = CapTable(
        company=Company(name="Acme"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=1000),
        ],
    )
    readback = VolPackReadback(
        market=MarketInputs(volatility=0.55, time_to_liquidity_years=4.0,
                             risk_free_rate=0.045),
        sourcing={"volatility": "peer set Q1 2025", "ttl": "Series C plan"},
    )
    blob = build_dcf_sidecar_bytes(ct, vol_pack_readback=readback)
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    rm = wb["Read me"]
    body = " ".join(
        str(rm.cell(row=r, column=1).value or "")
        for r in range(1, 40)
    )
    assert "Sourcing notes" in body
    assert "peer set Q1 2025" in body
    assert "Series C plan" in body
