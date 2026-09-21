"""W5.1 / W5.2 / W5.3 / W5.5 / W5.7 — HTMX UI surface tests."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

import openpyxl
import pytest
from flask import Flask

from src.cookie_auth import (
    COOKIE_NAME,
    attach_login_blueprint,
    issue_session_cookie,
    verify_session_cookie,
)
from src.engagement import EngagementStore
from src.engagement_routes import attach_engagement_blueprint
from src.identity import Role, StaticUserProvider, User
from src.rate_limit import ExportRateLimiter


@pytest.fixture
def web():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as eng_fh, \
         tempfile.NamedTemporaryFile(suffix=".db", delete=False) as ex_fh, \
         tempfile.NamedTemporaryFile(suffix=".db", delete=False) as cp_fh:
        eng_path = Path(eng_fh.name)
        ex_path = Path(ex_fh.name)
        cp_path = Path(cp_fh.name)
    store = EngagementStore(db_path=eng_path)
    export_limiter = ExportRateLimiter(db_path=ex_path, soft_limit=50, hard_limit=100)
    compute_limiter = ExportRateLimiter(db_path=cp_path, soft_limit=3, hard_limit=10)
    app = Flask(__name__)
    app.config["TESTING"] = True
    app.config["SESSION_SECRET_KEY"] = b"x" * 32
    users = {
        "tok-a": User(id="u-a", email="a@x", role=Role.analyst, display_name="A"),
        "tok-r": User(id="u-r", email="r@x", role=Role.reviewer, display_name="R"),
    }
    attach_engagement_blueprint(app, store, StaticUserProvider(users),
                                 export_limiter=export_limiter,
                                 compute_limiter=compute_limiter)
    attach_login_blueprint(app)
    yield app.test_client(), store
    eng_path.unlink(missing_ok=True)
    ex_path.unlink(missing_ok=True)
    cp_path.unlink(missing_ok=True)


def _h(tok: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {tok}"}


def _upload(client, eng_id, tok="tok-a"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 5_000_000, 0.01, "2020-01-01",
               "", "", "", ""])
    ws.append(["Series A", "preferred", 1_000_000, 1.00, "2024-01-01",
               1.0, "non_participating", 1, ""])  # blank AD → blocker
    buf = io.BytesIO()
    wb.save(buf)
    return client.post(
        f"/engagement/{eng_id}/upload", headers=_h(tok),
        data={"file": (io.BytesIO(buf.getvalue()), "demo.xlsx"),
              "expected_version": "0"},
        content_type="multipart/form-data",
    )


def _new_engagement(client, tok="tok-a"):
    return client.post(
        "/engagement/", headers=_h(tok),
        json={"client_id": "c", "valuation_date": "2025-06-01"},
    ).get_json()["id"]


# =============================================================================
# W5.1 — Snapshot-aware /whatif form generation
# =============================================================================


def test_w51_detail_lists_preferred_classes_for_whatif(web):
    client, _ = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    r = client.get(f"/engagement/{eng_id}?html=1", headers=_h("tok-a"))
    body = r.get_data(as_text=True)
    # The detail page should render an override row for "Series A"
    assert "Series A" in body
    assert 'name="shares_Series A"' in body
    assert 'name="lp_mult_Series A"' in body


def test_w51_detail_without_snapshot_shows_no_overrides(web):
    client, _ = web
    eng_id = _new_engagement(client)
    r = client.get(f"/engagement/{eng_id}?html=1", headers=_h("tok-a"))
    body = r.get_data(as_text=True)
    # No preferred class rows
    assert 'name="shares_Series A"' not in body
    assert "No preferred classes on the head snapshot yet" in body


# =============================================================================
# W5.2 — Inline finding-resolution form
# =============================================================================


def test_w52_finding_list_renders_inline(web):
    client, _ = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    r = client.get(f"/engagement/{eng_id}?html=1", headers=_h("tok-a"))
    body = r.get_data(as_text=True)
    # Series A has no AD → AD-MISSING blocker fires
    assert "AD-MISSING" in body
    # The inline Resolve form exists (analyst can resolve)
    assert "Record resolution" in body


def test_w52_resolve_form_records_resolution(web):
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    r = client.post(
        f"/engagement/{eng_id}/resolve?html=1", headers=_h("tok-a"),
        data={
            "snapshot_id": store.get_engagement(eng_id).head_snapshot_id,
            "finding_code": "AD-MISSING-Series A",
            "decision": '{"variant":"broad_based_weighted_average"}',
            "citation": "Charter §4.3(a)",
        },
    )
    assert r.status_code == 200
    assert "resolved at" in r.get_data(as_text=True)


def test_w52_resolve_form_requires_citation(web):
    """Empty citation must be rejected (spec §3.2 — resolution citation
    required). The decorator now returns HTML banner for HTMX path."""
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    r = client.post(
        f"/engagement/{eng_id}/resolve?html=1", headers=_h("tok-a"),
        data={
            "snapshot_id": store.get_engagement(eng_id).head_snapshot_id,
            "finding_code": "AD-MISSING-Series A",
            "decision": "{}",
            "citation": "",
        },
    )
    # M-3 fix: HTML branch on errors
    assert r.status_code in (400, 409)
    assert "<html" in r.get_data(as_text=True)


# =============================================================================
# W5.3 — Snapshot diff visualization
# =============================================================================


def test_w53_subsequent_events_panel_renders_after_two_snapshots(web):
    client, _ = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    # Upload a second snapshot with different share counts
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 6_000_000, 0.01, "2020-01-01",
               "", "", "", ""])  # changed!
    ws.append(["Series A", "preferred", 1_000_000, 1.00, "2024-01-01",
               1.0, "non_participating", 1, ""])
    buf = io.BytesIO()
    wb.save(buf)
    client.post(
        f"/engagement/{eng_id}/upload", headers=_h("tok-a"),
        data={"file": (io.BytesIO(buf.getvalue()), "demo2.xlsx"),
              "expected_version": "1"},
        content_type="multipart/form-data",
    )
    r = client.get(f"/engagement/{eng_id}?html=1", headers=_h("tok-a"))
    body = r.get_data(as_text=True)
    assert "Subsequent events" in body
    assert "Share count" in body


def test_w53_no_subsequent_events_when_single_snapshot(web):
    client, _ = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    r = client.get(f"/engagement/{eng_id}?html=1", headers=_h("tok-a"))
    body = r.get_data(as_text=True)
    assert "Subsequent events" not in body


# =============================================================================
# W5.5 — /whatif compute-budget rate limit
# =============================================================================


def test_w55_whatif_hits_compute_limit_separate_from_export(web):
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id)
    # Compute limiter soft=3 → 4th /whatif should 429
    for i in range(3):
        r = client.post(f"/engagement/{eng_id}/whatif?html=1",
                        headers=_h("tok-a"), data={})
        assert r.status_code == 200
    r = client.post(f"/engagement/{eng_id}/whatif?html=1",
                    headers=_h("tok-a"), data={})
    assert r.status_code == 429
    assert r.get_json()["error_code"] == "compute-rate-limit"


def test_w55_whatif_does_not_burn_export_budget(web):
    """Even after exhausting compute budget, the export budget should be
    untouched and the memo should still generate."""
    client, store = web
    eng_id = _new_engagement(client)
    # Need a clean cap table (no blockers) for memo to render
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 1_000_000, 0.01, "2020-01-01",
               "", "", "", ""])
    ws.append(["Series A", "preferred", 1_000_000, 1.00, "2024-01-01",
               1.0, "non_participating", 1, "broad_based_weighted_average"])
    buf = io.BytesIO()
    wb.save(buf)
    client.post(
        f"/engagement/{eng_id}/upload", headers=_h("tok-a"),
        data={"file": (io.BytesIO(buf.getvalue()), "demo.xlsx"),
              "expected_version": "0"},
        content_type="multipart/form-data",
    )
    # Burn compute budget
    for _ in range(20):
        client.post(f"/engagement/{eng_id}/whatif?html=1",
                    headers=_h("tok-a"), data={})
    # Memo (different budget) should still work
    r = client.get(f"/engagement/{eng_id}/memo.pdf?reviewer=R",
                    headers=_h("tok-r"))
    assert r.status_code != 429  # export budget untouched


# =============================================================================
# W5.7 — Cookie auth shim
# =============================================================================


def test_w57_login_get_returns_html_form():
    """The /login GET page renders the dev-stub login form."""
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
        p = Path(fh.name)
    try:
        store = EngagementStore(db_path=p)
        app = Flask(__name__)
        app.config["SESSION_SECRET_KEY"] = b"x" * 32
        users = {"tok-1": User(id="u-1", email="a@x", role=Role.analyst,
                               display_name="A")}
        attach_engagement_blueprint(app, store, StaticUserProvider(users))
        attach_login_blueprint(app)
        client = app.test_client()
        r = client.get("/login")
        assert r.status_code == 200
        body = r.get_data(as_text=True)
        assert "<form" in body
        assert 'name="token"' in body
    finally:
        p.unlink(missing_ok=True)


def test_w57_login_post_issues_cookie_and_redirects():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
        p = Path(fh.name)
    try:
        store = EngagementStore(db_path=p)
        app = Flask(__name__)
        app.config["TESTING"] = True
        app.config["SESSION_SECRET_KEY"] = b"x" * 32
        users = {"tok-1": User(id="u-1", email="a@x", role=Role.analyst,
                               display_name="A")}
        attach_engagement_blueprint(app, store, StaticUserProvider(users))
        attach_login_blueprint(app)
        client = app.test_client()
        r = client.post("/login", data={"token": "tok-1",
                                         "next": "/engagement/?html=1"})
        assert r.status_code == 302
        assert r.headers["Location"] == "/engagement/?html=1"
        # Cookie set
        set_cookie = r.headers.get("Set-Cookie", "")
        assert COOKIE_NAME in set_cookie
    finally:
        p.unlink(missing_ok=True)


def test_w57_login_post_rejects_bad_token():
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
        p = Path(fh.name)
    try:
        store = EngagementStore(db_path=p)
        app = Flask(__name__)
        app.config["TESTING"] = True
        app.config["SESSION_SECRET_KEY"] = b"x" * 32
        users = {"tok-1": User(id="u-1", email="a@x", role=Role.analyst,
                               display_name="A")}
        attach_engagement_blueprint(app, store, StaticUserProvider(users))
        attach_login_blueprint(app)
        client = app.test_client()
        r = client.post("/login", data={"token": "tok-WRONG"})
        assert r.status_code == 401
    finally:
        p.unlink(missing_ok=True)


def test_w57_cookie_authenticates_engagement_routes(web):
    """A request with the session cookie set should succeed without a
    Bearer header."""
    client, _ = web
    # First, log in to receive the cookie
    r = client.post("/login", data={"token": "tok-a", "next": "/engagement/?html=1"})
    assert r.status_code == 302
    # Now hit engagement list WITHOUT Authorization header — cookie auth
    # should resolve via the test_client's cookie jar
    r = client.get("/engagement/?html=1")
    assert r.status_code == 200


def test_w57_cookie_signature_verifies():
    secret = b"abc" * 16
    cookie = issue_session_cookie("u-42", secret, ttl_days=1)
    assert verify_session_cookie(cookie, secret) == "u-42"


def test_w57_tampered_cookie_rejected():
    secret = b"abc" * 16
    cookie = issue_session_cookie("u-42", secret, ttl_days=1)
    # Tamper with the payload
    parts = cookie.split(".")
    tampered = "BAD" + parts[0][3:] + "." + parts[1]
    assert verify_session_cookie(tampered, secret) is None


def test_w57_cookie_with_wrong_secret_rejected():
    cookie = issue_session_cookie("u-42", b"a" * 32, ttl_days=1)
    assert verify_session_cookie(cookie, b"b" * 32) is None


def test_w57_html_unauthenticated_redirects_to_login(web):
    """Browser-style request without cookie or Bearer should 302 to /login."""
    client, _ = web
    r = client.get("/engagement/?html=1", follow_redirects=False)
    assert r.status_code == 302
    assert r.headers["Location"].startswith("/login")


def test_w57_json_unauthenticated_still_401s(web):
    """API clients without Bearer still get 401, not a redirect."""
    client, _ = web
    r = client.get("/engagement/")
    assert r.status_code == 401
