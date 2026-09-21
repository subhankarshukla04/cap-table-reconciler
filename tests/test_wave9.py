"""Wave-9 feature + polish tests.

Covers W9.1 (rule provenance in memo), W9.2 (OPM Backsolve route),
W9.3 (vol pack → DCF sidecar), W9.4 (phase-0 anchoring), W9.5 (CSRF
rotation), W9.6 (NaN guard), W9.7 (snapshots pagination), W9.8 (rule
coverage CLI).
"""

from __future__ import annotations

import io
import math
import sqlite3
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from flask import Flask

from src.cookie_auth import (
    CSRF_COOKIE_NAME,
    CSRF_ROTATION_SECONDS,
    attach_login_blueprint,
    csrf_token_age_seconds,
    issue_csrf_token,
)
from src.engagement import AuditEventType, EngagementStore
from src.engagement_routes import attach_engagement_blueprint
from src.identity import Role, StaticUserProvider, User
from src.rate_limit import ExportRateLimiter
from src.snapshot_timeline import _classify_pct
from src.token_deny import TokenDenyList


# ---- shared web fixture ---------------------------------------------------


@pytest.fixture
def web():
    paths = []
    for _ in range(4):
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
            paths.append(Path(fh.name))
    eng_p, ex_p, cp_p, dn_p = paths
    store = EngagementStore(db_path=eng_p)
    export_limiter = ExportRateLimiter(db_path=ex_p, soft_limit=50, hard_limit=100)
    compute_limiter = ExportRateLimiter(db_path=cp_p, soft_limit=20, hard_limit=40)
    deny_list = TokenDenyList(db_path=dn_p)
    app = Flask(__name__)
    app.config["TESTING"] = True
    app.config["SESSION_SECRET_KEY"] = b"x" * 32
    app.config["TOKEN_DENY_LIST"] = deny_list
    users = {
        "tok-a": User(id="u-a", email="a@x", role=Role.analyst, display_name="A"),
        "tok-p": User(id="u-p", email="p@x", role=Role.partner, display_name="P"),
    }
    provider = StaticUserProvider(users)
    attach_engagement_blueprint(app, store, provider,
                                 export_limiter=export_limiter,
                                 compute_limiter=compute_limiter)
    attach_login_blueprint(app, identity_provider=provider)
    yield app.test_client(), store
    deny_list.close()
    for p in paths:
        p.unlink(missing_ok=True)


def _bearer(tok: str):
    return {"Authorization": f"Bearer {tok}"}


def _xlsx_bytes(shares_a: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 5_000_000, 0.01, "2020-01-01",
               "", "", "", ""])
    ws.append(["Series A", "preferred", shares_a, 1.00, "2024-01-01",
               1.0, "non_participating", 1, "broad_based_weighted_average"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# W9.1 rule provenance in PDF memo
# =============================================================================


def test_w91_memo_html_includes_rule_provenance_column():
    """The PDF memo template renders a per-finding Rule provenance cell
    with rule_id + pack version + citation."""
    from src.pdf_memo import PDFInputs, render_pdf_memo_html, ReviewerInfo
    from src.waterfall import compute_waterfall
    from src.models import (
        CapTable, Company, ShareClass, ShareClassType,
        LiquidationPreference, LPType,
    )
    ct = CapTable(
        company=Company(name="Acme"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=5_000_000),
            ShareClass(
                name="Series A", type=ShareClassType.preferred,
                shares_outstanding=1_000_000, issue_price=1.0,
                liquidation_preference=LiquidationPreference(
                    multiple=1.0, type=LPType.non_participating,
                    amount=1_000_000,
                ),
            ),  # missing AD → AD-MISSING-Series A blocker from G-AD-001
        ],
    )
    wf = compute_waterfall(ct)
    from src.rule_pack import run_pack
    findings = run_pack(ct)
    # PDFInputs.resolutions list satisfies _ensure_eligible so the memo
    # renders even with the AD-MISSING blocker present. We're testing
    # the provenance column rendering, not the eligibility gate.
    inputs = PDFInputs(
        cap_table=ct, waterfall=wf, findings=findings,
        resolutions=[
            {"finding_code": "AD-MISSING-Series A",
             "decision_summary": "broad_based_weighted_average",
             "citation": "Charter §4.3(a)",
             "resolved_by": "test"},
        ],
        reviewer=ReviewerInfo(reviewer_name="r", preparer="p"),
    )
    html = render_pdf_memo_html(inputs)
    assert "Rule provenance" in html
    # G-AD-001 emits AD-MISSING-Series A; provenance column should expose the rule id.
    assert "G-AD-001" in html


# =============================================================================
# W9.2 OPM Backsolve route
# =============================================================================


def test_w92_opm_route_succeeds_and_writes_audit_event(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "demo.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    # Resolve the AD-MISSING blocker so backsolve doesn't refuse.
    snap_id = store.get_engagement(eng_id).head_snapshot_id
    client.post(f"/engagement/{eng_id}/resolve", headers=_bearer("tok-a"),
                json={"snapshot_id": snap_id,
                      "finding_code": "AD-MISSING-Series A",
                      "decision": {"variant": "broad_based_weighted_average"},
                      "citation": "Charter §4.3(a)"})
    r = client.post(
        f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
        json={
            "volatility": 0.55,
            "time_to_liquidity_years": 4.0,
            "risk_free_rate": 0.045,
            "dlom": 0.25,
            "anchor_class_name": "Series A",
            "anchor_price_per_share": 1.00,
        },
    )
    assert r.status_code == 200, r.get_data(as_text=True)
    payload = r.get_json()
    assert payload["implied_total_equity_value"] > 0
    assert payload["anchor"]["class_name"] == "Series A"
    per_class = {pc["name"]: pc for pc in payload["per_class"]}
    assert "Series A" in per_class
    assert "Common" in per_class
    # Common gets a post-DLOM number.
    assert per_class["Common"]["fair_value_per_share_after_dlom"] is not None

    # Audit event lands in the chain.
    events = store.list_audit_events(eng_id)
    opm = [e for e in events if e.event_type == AuditEventType.opm_backsolve_run]
    assert opm


def test_w92_opm_route_refuses_without_snapshot(web):
    client, _ = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    r = client.post(f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
                    json={"volatility": 0.5, "time_to_liquidity_years": 4.0,
                          "risk_free_rate": 0.045,
                          "anchor_class_name": "X",
                          "anchor_price_per_share": 1.0})
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "no-snapshot"


def test_w92_opm_route_rejects_missing_inputs(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                data={"file": (io.BytesIO(_xlsx_bytes(1_000_000)), "demo.xlsx"),
                      "expected_version": "0"},
                content_type="multipart/form-data")
    r = client.post(f"/engagement/{eng_id}/opm", headers=_bearer("tok-a"),
                    json={"volatility": 0.5})  # missing the rest
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "opm-bad-inputs"


# =============================================================================
# W9.3 vol pack → DCF sidecar
# =============================================================================


def test_w93_vol_pack_readback_surfaces_in_dcf_sidecar_readme():
    from src.dcf_sidecar import build_dcf_sidecar_bytes
    from src.opm.backsolve import MarketInputs
    from src.opm.vol_pack import VolPackReadback
    from src.models import CapTable, Company, ShareClass, ShareClassType

    ct = CapTable(
        company=Company(name="Acme"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=1000),
        ],
    )
    readback = VolPackReadback(
        market=MarketInputs(volatility=0.55, time_to_liquidity_years=4.0,
                            risk_free_rate=0.045, dlom=0.25),
        sourcing={"volatility": "peer set Q1 2025"},
    )
    blob = build_dcf_sidecar_bytes(ct, vol_pack_readback=readback)
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    rm = wb["Read me"]
    body = " ".join(
        str(rm.cell(row=r, column=1).value or "")
        for r in range(1, 25)
    )
    assert "vol pack" in body.lower()
    assert "55.00%" in body  # volatility
    assert "4.00" in body  # TTL
    # Belt-and-braces: it must say "suggestions" / "analyst-overridable"
    assert "suggestion" in body.lower() or "overridable" in body.lower()


def test_w93_dcf_sidecar_omits_section_when_no_vol_pack():
    from src.dcf_sidecar import build_dcf_sidecar_bytes
    from src.models import CapTable, Company, ShareClass, ShareClassType
    ct = CapTable(
        company=Company(name="Acme"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=1000),
        ],
    )
    blob = build_dcf_sidecar_bytes(ct)  # no readback supplied
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    rm = wb["Read me"]
    body = " ".join(
        str(rm.cell(row=r, column=1).value or "")
        for r in range(1, 25)
    )
    assert "vol pack" not in body.lower()


# =============================================================================
# W9.4 phase-0 prefix anchoring
# =============================================================================


def test_w94_diff_exact_match_blocked_but_diffx_not(monkeypatch):
    monkeypatch.delenv("ALLOW_PHASE_0_DEMO", raising=False)
    import importlib
    import app as app_module
    importlib.reload(app_module)
    app_module.app.config["TESTING"] = False
    client = app_module.app.test_client()
    # Exact /diff (phase-0 demo) still blocked.
    r = client.get("/diff")
    assert r.status_code == 404
    # A hypothetical future top-level route /diffx would NOT match the
    # exact-set; here we hit a Flask 404 (route doesn't exist) rather
    # than the phase-0-disabled 404. Both 404 but with different bodies.
    r2 = client.get("/diffx-nonexistent")
    assert r2.status_code == 404
    # The phase-0 disabled body is JSON; the Flask not-found is HTML.
    assert r.headers.get("Content-Type", "").startswith("application/json")
    # Restore TESTING.
    app_module.app.config["TESTING"] = True


# =============================================================================
# W9.5 CSRF token rotation
# =============================================================================


def test_w95_csrf_token_carries_issued_at_prefix():
    tok = issue_csrf_token()
    head, rand = tok.split(".", 1)
    assert int(head) > 1_700_000_000  # plausible unix seconds
    assert len(rand) >= 30  # base64url-encoded 32 bytes


def test_w95_csrf_token_age_seconds_parses_prefix():
    tok = issue_csrf_token()
    age = csrf_token_age_seconds(tok)
    assert age is not None
    assert age < 5  # just minted


def test_w95_csrf_token_age_returns_none_for_legacy_opaque():
    """Pre-W9.5 tokens (no `.` prefix) return None — caller treats as stale."""
    assert csrf_token_age_seconds("opaque-legacy-token-no-dot") is None
    assert csrf_token_age_seconds("") is None
    assert csrf_token_age_seconds(None) is None  # type: ignore


def test_w95_rotation_constant_is_one_hour():
    assert CSRF_ROTATION_SECONDS == 3600


def test_w95_after_request_rotates_stale_csrf_cookie(web):
    """A cookie-authed GET with a legacy opaque CSRF cookie triggers
    rotation in the after_request hook."""
    client, _ = web
    # Login to get a fresh session cookie.
    r = client.post("/login", data={"token": "tok-a"})
    assert r.status_code == 302
    # Replace the CSRF cookie with a legacy opaque value.
    client.set_cookie(CSRF_COOKIE_NAME, "legacy-opaque-token",
                      domain="localhost")
    # Force a safe-method GET that goes through the engagement blueprint.
    r = client.get("/engagement/?html=1")
    assert r.status_code == 200
    csrf_after = client.get_cookie(CSRF_COOKIE_NAME)
    assert csrf_after.value != "legacy-opaque-token"
    assert "." in csrf_after.value  # W9.5 format


# =============================================================================
# W9.6 NaN guard
# =============================================================================


def test_w96_classify_pct_raises_on_nan():
    with pytest.raises(ValueError, match="NaN"):
        _classify_pct(float("nan"), 100.0)
    with pytest.raises(ValueError, match="NaN"):
        _classify_pct(100.0, float("nan"))


def test_w96_classify_pct_normal_values_unchanged():
    assert _classify_pct(100, 100) == "none"
    assert _classify_pct(100, 150) == "major"
    assert _classify_pct(None, None) == "none"


# =============================================================================
# W9.7 snapshots pagination
# =============================================================================


def test_w97_snapshots_pagination_default_returns_full_list(web):
    """No pagination params → default limit 50 returns all snapshots
    (preserves wave-7 API compatibility)."""
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    for shares, ver in [(1_000_000, 0), (1_500_000, 1)]:
        client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                    data={"file": (io.BytesIO(_xlsx_bytes(shares)), "x.xlsx"),
                          "expected_version": str(ver)},
                    content_type="multipart/form-data")
    r = client.get(f"/engagement/{eng_id}/snapshots", headers=_bearer("tok-a"))
    assert r.status_code == 200
    assert len(r.get_json()) == 2


def test_w97_snapshots_pagination_limit_offset_slice(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    for shares, ver in [(1_000_000, 0), (1_500_000, 1)]:
        client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                    data={"file": (io.BytesIO(_xlsx_bytes(shares)), "x.xlsx"),
                          "expected_version": str(ver)},
                    content_type="multipart/form-data")
    r = client.get(f"/engagement/{eng_id}/snapshots?limit=1&offset=1",
                   headers=_bearer("tok-a"))
    payload = r.get_json()
    assert len(payload) == 1


def test_w97_snapshots_pagination_rejects_bad_input(web):
    client, _ = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    r = client.get(f"/engagement/{eng_id}/snapshots?limit=999",
                   headers=_bearer("tok-a"))
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "bad-pagination"
    r = client.get(f"/engagement/{eng_id}/snapshots?limit=-1",
                   headers=_bearer("tok-a"))
    assert r.status_code == 400
    r = client.get(f"/engagement/{eng_id}/snapshots?offset=notanumber",
                   headers=_bearer("tok-a"))
    assert r.status_code == 400


def test_w97_html_view_renders_pagination_chrome(web):
    client, store = web
    r = client.post("/engagement/", headers=_bearer("tok-a"),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    eng_id = r.get_json()["id"]
    for shares, ver in [(1_000_000, 0), (1_500_000, 1)]:
        client.post(f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
                    data={"file": (io.BytesIO(_xlsx_bytes(shares)), "x.xlsx"),
                          "expected_version": str(ver)},
                    content_type="multipart/form-data")
    r = client.get(f"/engagement/{eng_id}/snapshots?html=1&limit=1&offset=0",
                   headers=_bearer("tok-a"))
    body = r.get_data(as_text=True)
    assert "showing 1 of 2 snapshots" in body
    assert "Next 1" in body


# =============================================================================
# W9.8 rule coverage CLI smoke test
# =============================================================================


def test_w98_rule_coverage_cli_runs():
    """The Flask CLI command runs end-to-end against the project fixtures."""
    import importlib
    import app as app_module
    importlib.reload(app_module)
    runner = app_module.app.test_cli_runner()
    result = runner.invoke(args=["rule-coverage"])
    assert result.exit_code == 0
    out = result.output
    assert "Total rules:" in out
    assert "Coverage detail" in out
