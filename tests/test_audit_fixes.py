"""Regression tests for the pre-build hardening pass (CODE_AUDIT.md).

One test per audit finding fixed. Each test reproduces the original bug
condition and asserts the fixed behaviour. These tests are the load-bearing
discipline: if a future refactor reintroduces a regression, one of these
fails first.
"""

from __future__ import annotations

import gc
import tempfile
import warnings
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from src.audit_memo import build_audit_memo
from src.checklist import run_checklist
from src.models import (
    CapTable,
    Company,
    LiquidationPreference,
    LPType,
    SAFE,
    ShareClass,
    ShareClassType,
)
from src.parser import _detect_field_for_header, load_from_canonical_json, parse_excel
from src.persistence import SessionStore
from src.waterfall import compute_waterfall


# ---- BUG-001 — unknown class type must not silently drop the row ------------


def test_bug001_unknown_class_type_preserves_shares():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date"])
    ws.append(["Founders Stock", "Founders", 100_000, 0.001, "2020-01-01"])
    ws.append(["Common Stock", "Common", 50_000, 0.01, "2020-01-02"])
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        wb.save(fh.name)
        path = Path(fh.name)
    try:
        ct, report = parse_excel(path)
    finally:
        path.unlink(missing_ok=True)
    names = {sc.name for sc in ct.share_classes}
    assert "Founders Stock" in names, "row with unknown type 'Founders' was silently dropped"
    assert "Common Stock" in names
    assert any(w.code == "class_type_unknown" for w in report.warnings)


# ---- BUG-002 — header detection must be position-aware -----------------------


@pytest.mark.parametrize(
    "header, expected",
    [
        ("share class type", "class_type"),
        ("Share Class Type", "class_type"),
        ("Class Name", "class_name"),
        ("Type", "class_type"),
        ("Comments on shares", "notes"),
        ("Notes regarding voting", "notes"),
        ("Shareholder Name", "class_name"),
        ("Outstanding shares", "shares"),
    ],
)
def test_bug002_header_field_detection(header, expected):
    assert _detect_field_for_header(header) == expected


# ---- BUG-003 — SAFE-UNCONVERTED fires on any qualifying subsequent round ----


def test_bug003_safe_unconverted_checks_all_subsequent_rounds():
    ct = CapTable(
        company=Company(name="T"),
        share_classes=[
            ShareClass(name="C", type=ShareClassType.common, shares_outstanding=1000),
            ShareClass(
                name="A",
                type=ShareClassType.preferred,
                shares_outstanding=1000,
                issue_price=0.5,
                issue_date=date(2022, 1, 1),
                seniority_rank=2,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=500, type=LPType.non_participating
                ),
            ),
            ShareClass(
                name="B",
                type=ShareClassType.preferred,
                shares_outstanding=10_000,
                issue_price=2.0,
                issue_date=date(2023, 1, 1),
                seniority_rank=1,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=20_000, type=LPType.non_participating
                ),
            ),
        ],
        safes_outstanding=[
            SAFE(
                id="SAFE-1",
                principal=5_000,
                issue_date=date(2021, 6, 1),
                conversion_trigger_threshold=10_000,
            )
        ],
    )
    safe_findings = [f for f in run_checklist(ct) if "SAFE-UNCONVERTED" in f.code]
    assert len(safe_findings) == 1
    assert safe_findings[0].severity == "blocker"


# ---- BUG-004 — audit memo escapes pipe characters in every cell -------------


def test_bug004_audit_memo_escapes_pipe_characters():
    ct = CapTable(
        company=Company(name="Acme | Co"),
        share_classes=[
            ShareClass(
                name="Common|Stock",
                type=ShareClassType.common,
                shares_outstanding=1000,
            ),
            ShareClass(
                name="A",
                type=ShareClassType.preferred,
                shares_outstanding=100,
                seniority_rank=1,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=100, type=LPType.non_participating
                ),
            ),
        ],
    )
    memo = build_audit_memo(ct, compute_waterfall(ct), run_checklist(ct))
    assert "Acme \\| Co" in memo
    assert "Common\\|Stock" in memo
    # And the un-escaped form must NOT appear inside a table cell context.
    assert "| Common|Stock |" not in memo


# ---- BUG-005 — SessionStore must not leak SQLite connections ---------------


def test_bug005_session_store_no_resource_warnings():
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always", ResourceWarning)
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
            db_path = Path(fh.name)
        try:
            store = SessionStore(db_path)
            ct = CapTable(
                company=Company(name="X"),
                share_classes=[
                    ShareClass(
                        name="C", type=ShareClassType.common, shares_outstanding=1
                    )
                ],
            )
            store["t1"] = {"cap_table": ct, "parse_report": None}
            assert len(store) == 1
            assert "t1" in store
            assert store.get("t1") is not None
            assert store.delete("t1")
            store.clear()
            del store
            gc.collect()
        finally:
            db_path.unlink(missing_ok=True)
    leaks = [w for w in caught if issubclass(w.category, ResourceWarning)]
    assert leaks == [], f"SessionStore leaked: {[str(w.message) for w in leaks]}"


# ---- BUG-007 — /healthz returns only {"status": "ok"} -----------------------


def test_bug007_healthz_matches_spec():
    from app import app

    client = app.test_client()
    r = client.get("/healthz")
    assert r.status_code == 200
    body = r.get_json()
    assert body == {"status": "ok"}


# ---- BUG-009 — load_from_canonical_json catches inconsistent cap fields -----


def test_bug009_inconsistent_cap_raises():
    bad = {
        "company": {"name": "T", "currency": "USD"},
        "share_classes": [
            {
                "name": "C",
                "type": "common",
                "shares_outstanding": 100,
                "seniority_rank": 99,
                "liquidation_preference": None,
                "anti_dilution": None,
            },
            {
                "name": "A",
                "type": "preferred",
                "shares_outstanding": 100,
                "issue_price_usd": 1.0,
                "issue_date": "2020-01-01",
                "seniority_rank": 1,
                "liquidation_preference": {
                    "multiple": 1.0,
                    "amount_usd": 100.0,
                    "type": "participating_capped",
                    "cap_multiple": 2.0,
                    "cap_amount_usd": 1_000_000.0,
                },
                "anti_dilution": None,
                "participation": {"mode": "with_cap", "cap_multiple_of_lp": 2.0},
            },
        ],
    }
    import json as _json

    with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w") as fh:
        fh.write(_json.dumps(bad))
        p = Path(fh.name)
    try:
        with pytest.raises(ValueError, match="Inconsistent cap"):
            load_from_canonical_json(p)
    finally:
        p.unlink(missing_ok=True)


# ---- BUG-013 — Company tab tolerates real-world label variants --------------


def test_bug013_company_tab_aliased_labels():
    wb = openpyxl.Workbook()
    co = wb.create_sheet("Company")
    co.append(["Company Name", "BugThirteen Inc"])
    co.append(["Country of Incorporation", "Singapore"])
    co.append(["Date of Valuation", "2025-12-31"])
    co.append(["Currency Code", "SGD"])
    ws = wb.create_sheet("Cap Table")
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date"])
    ws.append(["Common", "common", 1000, 0.0001, "2020-01-01"])
    # Remove the default Sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        wb.save(fh.name)
        p = Path(fh.name)
    try:
        ct, _ = parse_excel(p)
    finally:
        p.unlink(missing_ok=True)
    assert ct.company.name == "BugThirteen Inc"
    assert ct.company.jurisdiction == "Singapore"
    assert ct.company.currency == "SGD"
    assert ct.company.valuation_date == date(2025, 12, 31)
