"""DCF input sidecar tests (SYSTEM_SPEC §5.3)."""

from __future__ import annotations

import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.dcf_sidecar import _slug_for_defined_name, build_dcf_sidecar, build_dcf_sidecar_bytes
from src.models import (
    AntiDilution,
    AntiDilutionVariant,
    CapTable,
    Company,
    LiquidationPreference,
    LPType,
    ShareClass,
    ShareClassType,
)


def _ct():
    return CapTable(
        company=Company(name="DCF Co", currency="USD"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common, shares_outstanding=5_000_000),
            ShareClass(
                name="Series A & B",  # ampersand requires sanitisation
                type=ShareClassType.preferred,
                shares_outstanding=1_000_000,
                issue_price=1.0,
                issue_date=date(2024, 1, 1),
                seniority_rank=1,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=1_000_000, type=LPType.non_participating
                ),
                anti_dilution=AntiDilution(variant=AntiDilutionVariant.broad_based_weighted_average),
            ),
        ],
    )


def test_slug_sanitises_invalid_chars():
    assert _slug_for_defined_name("Series A & B") == "Series_A___B"
    assert _slug_for_defined_name("1stClass") == "_1stClass"
    assert _slug_for_defined_name("C") == "_C"  # reserved


def test_workbook_has_defined_names_per_class():
    wb = build_dcf_sidecar(_ct())
    name_keys = {n.name for n in wb.defined_names.values()}
    assert "share_count_Common" in name_keys
    assert "lp_amount_Common" in name_keys
    assert "conv_ratio_Common" in name_keys
    assert "share_count_Series_A___B" in name_keys
    assert "lp_amount_Series_A___B" in name_keys
    assert "total_fully_diluted" in name_keys
    assert "lp_total" in name_keys


def test_workbook_round_trips_via_openpyxl():
    """The bytes must be a valid xlsx that openpyxl reopens cleanly with
    all defined names intact."""
    raw = build_dcf_sidecar_bytes(_ct())
    wb = openpyxl.load_workbook(BytesIO(raw))
    assert "Cap Inputs" in wb.sheetnames
    assert "Named Range Map" in wb.sheetnames
    assert "Read me" in wb.sheetnames
    # Cap-table values land in the expected cells.
    sheet = wb["Cap Inputs"]
    assert sheet["C2"].value == 5_000_000  # Common share count
    assert sheet["C3"].value == 1_000_000  # Series share count
    assert sheet["D3"].value == 1_000_000  # Series LP amount


def test_named_range_map_preserves_display_name():
    wb = build_dcf_sidecar(_ct())
    ns = wb["Named Range Map"]
    display_names = [ns.cell(row=r, column=1).value for r in range(2, 4)]
    assert "Series A & B" in display_names
