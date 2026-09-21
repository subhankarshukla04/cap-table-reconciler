"""Wave-7 snapshot timeline + N-way diff tests.

Coverage:
  - snapshot_timeline.compute_timeline_diff() — magnitude classification,
    presence handling, categorical vs numeric fields.
  - engagement_routes /snapshots + /diff + /diff.xlsx — end-to-end.
  - audit-log integration: snapshot_diff_exported event lands on xlsx
    download.
  - W7.4 change_note round-trips through upload → snapshot → diff header.
"""

from __future__ import annotations

import io
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pytest
from flask import Flask

from src.cookie_auth import attach_login_blueprint
from src.diff_workpaper import build_diff_workpaper_xlsx
from src.engagement import AuditEventType, EngagementStore
from src.engagement_routes import attach_engagement_blueprint
from src.identity import Role, StaticUserProvider, User
from src.models import (
    CapTable, Company, ShareClass, ShareClassType,
    LiquidationPreference, LPType, AntiDilution, AntiDilutionVariant,
)
from src.rate_limit import ExportRateLimiter
from src.snapshot_timeline import (
    SnapshotMeta,
    _classify_pct,
    _max_tier,
    compute_timeline_diff,
)
from src.token_deny import TokenDenyList


# ---- Pure compute ---------------------------------------------------------


@pytest.mark.parametrize("prior,current,expected", [
    (None, None, "none"),
    (None, 100, "major"),
    (100, None, "major"),
    (0, 0, "none"),
    (0, 5, "major"),
    (100, 100, "none"),
    (100, 104, "minor"),     # 4% drift
    (100, 110, "material"),  # 10% drift
    (100, 150, "major"),     # 50% drift
])
def test_classify_pct(prior, current, expected):
    assert _classify_pct(prior, current) == expected


def test_max_tier_picks_strictest():
    assert _max_tier(["none", "minor", "material"]) == "material"
    assert _max_tier(["minor", "major"]) == "major"
    assert _max_tier([]) == "none"


def _mk(shares_a: int, lp_amount: float = None, variant=LPType.non_participating):
    return CapTable(
        company=Company(name="Acme"),
        share_classes=[
            ShareClass(name="Common", type=ShareClassType.common,
                       shares_outstanding=5_000_000),
            ShareClass(
                name="Series A", type=ShareClassType.preferred,
                shares_outstanding=shares_a, issue_price=1.0,
                liquidation_preference=LiquidationPreference(
                    multiple=1.0, type=variant,
                    amount=lp_amount if lp_amount is not None else shares_a,
                ),
            ),
        ],
    )


def _metas(n: int) -> list[SnapshotMeta]:
    return [
        SnapshotMeta(id=f"s{i}", created_at=datetime.now(timezone.utc),
                      source_filename=f"f{i}.xlsx", created_by="u")
        for i in range(n)
    ]


def test_compute_timeline_diff_flags_major_share_jump():
    diff = compute_timeline_diff(_metas(3),
                                  [_mk(1_000_000), _mk(1_050_000), _mk(1_500_000)])
    series_a = next(r for r in diff.class_rows if r.class_name == "Series A")
    assert series_a.overall_tier == "major"
    assert series_a.field_tiers["shares_outstanding"] == "major"


def test_compute_timeline_diff_categorical_lp_variant_change_is_major():
    diff = compute_timeline_diff(
        _metas(2),
        [_mk(1_000_000, variant=LPType.non_participating),
         _mk(1_000_000, variant=LPType.participating_uncapped)],
    )
    series_a = next(r for r in diff.class_rows if r.class_name == "Series A")
    assert series_a.overall_tier == "major"
    assert series_a.field_tiers["lp_variant_label"] == "major"


def test_compute_timeline_diff_no_drift_is_none():
    diff = compute_timeline_diff(_metas(2),
                                  [_mk(1_000_000), _mk(1_000_000)])
    assert all(r.overall_tier == "none" for r in diff.class_rows)
    assert diff.classes_changed == 0


def test_compute_timeline_diff_added_class_counts():
    ct1 = _mk(1_000_000)
    ct2 = _mk(1_000_000)
    ct2.share_classes.append(ShareClass(
        name="Series B", type=ShareClassType.preferred,
        shares_outstanding=2_000_000, issue_price=2.0,
        liquidation_preference=LiquidationPreference(
            multiple=1.0, type=LPType.non_participating, amount=4_000_000,
        ),
    ))
    diff = compute_timeline_diff(_metas(2), [ct1, ct2])
    assert diff.classes_added == 1
    series_b = next(r for r in diff.class_rows if r.class_name == "Series B")
    assert series_b.field_tiers["presence"] == "major"


def test_compute_timeline_diff_redacted_snapshots_render_as_blanks():
    diff = compute_timeline_diff(_metas(3),
                                  [_mk(1_000_000), None, _mk(1_500_000)])
    series_a = next(r for r in diff.class_rows if r.class_name == "Series A")
    # Middle column is None for the redacted snapshot.
    assert series_a.views[1].class_name is None


def test_compute_timeline_diff_refuses_single_snapshot():
    with pytest.raises(ValueError, match="at least 2"):
        compute_timeline_diff(_metas(1), [_mk(1)])


def test_compute_timeline_diff_length_mismatch_raises():
    with pytest.raises(ValueError):
        compute_timeline_diff(_metas(2), [_mk(1)])


# ---- Workpaper xlsx -------------------------------------------------------


def test_build_diff_workpaper_xlsx_has_expected_tabs():
    diff = compute_timeline_diff(
        _metas(2), [_mk(1_000_000), _mk(1_500_000)],
    )
    blob = build_diff_workpaper_xlsx(diff)
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    assert set(wb.sheetnames) == {"Summary", "Class Drift", "Snapshots"}


def test_build_diff_workpaper_xlsx_drift_tab_includes_magnitude():
    diff = compute_timeline_diff(
        _metas(2), [_mk(1_000_000), _mk(1_500_000)],
    )
    blob = build_diff_workpaper_xlsx(diff)
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    ws = wb["Class Drift"]
    # Find a row with field=shares_outstanding for Series A — it should
    # carry "major" in column C.
    found = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "Series A" and row[1] == "shares_outstanding":
            assert row[2] == "major"
            found = True
            break
    assert found, "no Series A shares_outstanding row in Class Drift tab"


# ---- End-to-end via Flask test client -------------------------------------


@pytest.fixture
def web():
    paths = []
    for _ in range(4):
        with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as fh:
            paths.append(Path(fh.name))
    eng_p, ex_p, cp_p, dn_p = paths
    store = EngagementStore(db_path=eng_p)
    export_limiter = ExportRateLimiter(db_path=ex_p, soft_limit=50, hard_limit=100)
    compute_limiter = ExportRateLimiter(db_path=cp_p, soft_limit=20, hard_limit=40)
    deny_list = TokenDenyList(db_path=dn_p)
    app = Flask(__name__)
    app.config["TESTING"] = True
    app.config["SESSION_SECRET_KEY"] = b"x" * 32
    app.config["TOKEN_DENY_LIST"] = deny_list
    users = {
        "tok-a": User(id="u-a", email="a@x", role=Role.analyst, display_name="A"),
    }
    provider = StaticUserProvider(users)
    attach_engagement_blueprint(app, store, provider,
                                 export_limiter=export_limiter,
                                 compute_limiter=compute_limiter)
    attach_login_blueprint(app, identity_provider=provider)
    yield app.test_client(), store
    deny_list.close()
    for p in paths:
        p.unlink(missing_ok=True)


def _bearer(tok: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {tok}"}


def _xlsx_bytes(shares_a: int) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cap Table"
    ws.append(["Class Name", "Type", "Shares", "Issue Price", "Issue Date",
               "LP Multiple", "LP Type", "Seniority", "Anti-Dilution"])
    ws.append(["Common", "common", 5_000_000, 0.01, "2020-01-01",
               "", "", "", ""])
    ws.append(["Series A", "preferred", shares_a, 1.00, "2024-01-01",
               1.0, "non_participating", 1, "broad_based_weighted_average"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _new_engagement(client, tok="tok-a"):
    r = client.post("/engagement/", headers=_bearer(tok),
                    json={"client_id": "c", "valuation_date": "2025-06-01"})
    return r.get_json()["id"]


def _upload(client, eng_id, shares_a, change_note=None, expected_version=0):
    data = {
        "file": (io.BytesIO(_xlsx_bytes(shares_a)), "demo.xlsx"),
        "expected_version": str(expected_version),
    }
    if change_note:
        data["change_note"] = change_note
    return client.post(
        f"/engagement/{eng_id}/upload", headers=_bearer("tok-a"),
        data=data, content_type="multipart/form-data",
    )


def test_timeline_route_lists_snapshots_chronologically(web):
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000, change_note="initial load",
            expected_version=0)
    _upload(client, eng_id, 1_500_000, change_note="series C closing",
            expected_version=1)
    r = client.get(f"/engagement/{eng_id}/snapshots", headers=_bearer("tok-a"))
    assert r.status_code == 200
    payload = r.get_json()
    assert len(payload) == 2
    assert payload[0]["created_at"] < payload[1]["created_at"]
    assert payload[0]["change_note"] == "initial load"
    assert payload[1]["change_note"] == "series C closing"
    assert payload[1]["is_head"] is True


def test_diff_route_refuses_single_snapshot(web):
    client, _ = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000, expected_version=0)
    r = client.get(f"/engagement/{eng_id}/diff?snap=any", headers=_bearer("tok-a"))
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "diff-needs-two-snapshots"


def test_diff_route_rejects_cross_engagement_snapshot(web):
    client, store = web
    eng_a = _new_engagement(client)
    eng_b = _new_engagement(client)
    _upload(client, eng_a, 1_000_000, expected_version=0)
    _upload(client, eng_b, 2_000_000, expected_version=0)
    snap_a = store.list_snapshots(eng_a)[0]
    snap_b = store.list_snapshots(eng_b)[0]
    # Try to diff snap from eng_a against snap from eng_b inside eng_a's URL.
    r = client.get(
        f"/engagement/{eng_a}/diff?snap={snap_a.id}&snap={snap_b.id}",
        headers=_bearer("tok-a"),
    )
    assert r.status_code == 400
    assert r.get_json()["error_code"] == "snapshot-engagement-mismatch"


def test_diff_route_returns_n_way_grid(web):
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000, expected_version=0)
    _upload(client, eng_id, 1_050_000, expected_version=1)  # minor
    _upload(client, eng_id, 1_500_000, expected_version=2)  # major
    snap_ids = [s.id for s in store.list_snapshots(eng_id)]
    qs = "&".join(f"snap={sid}" for sid in snap_ids)
    r = client.get(f"/engagement/{eng_id}/diff?{qs}", headers=_bearer("tok-a"))
    assert r.status_code == 200
    payload = r.get_json()
    assert len(payload["snapshots"]) == 3
    series_a_row = next(r for r in payload["rows"] if r["class_name"] == "Series A")
    assert series_a_row["overall_tier"] == "major"
    assert len(series_a_row["values"]) == 3
    assert series_a_row["values"][2]["shares_outstanding"] == 1_500_000


def test_diff_xlsx_export_writes_audit_event(web):
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000, expected_version=0)
    _upload(client, eng_id, 1_500_000, expected_version=1)
    snap_ids = [s.id for s in store.list_snapshots(eng_id)]
    qs = "&".join(f"snap={sid}" for sid in snap_ids)
    r = client.get(f"/engagement/{eng_id}/diff.xlsx?{qs}", headers=_bearer("tok-a"))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("Content-Type", "")
    # Verify the audit chain captured the export.
    events = store.list_audit_events(eng_id)
    assert any(e.event_type == AuditEventType.snapshot_diff_exported for e in events)


def test_timeline_html_renders_compare_checkbox_form(web):
    """The HTML view renders the compare form when >=2 snapshots exist."""
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000, expected_version=0)
    _upload(client, eng_id, 1_500_000, expected_version=1)
    r = client.get(f"/engagement/{eng_id}/snapshots?html=1",
                   headers=_bearer("tok-a"))
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "Compare:" in body
    assert 'type="checkbox" name="snap"' in body
    assert "View diff" in body


def test_change_note_round_trips_from_upload_to_diff_header(web):
    """W7.4: the change_note set at upload must surface in the diff JSON
    snapshot headers."""
    client, store = web
    eng_id = _new_engagement(client)
    _upload(client, eng_id, 1_000_000,
            change_note="initial cap from charter", expected_version=0)
    _upload(client, eng_id, 1_500_000,
            change_note="Series C closing 2025-06-01", expected_version=1)
    snap_ids = [s.id for s in store.list_snapshots(eng_id)]
    qs = "&".join(f"snap={sid}" for sid in snap_ids)
    r = client.get(f"/engagement/{eng_id}/diff?{qs}", headers=_bearer("tok-a"))
    notes = [s["change_note"] for s in r.get_json()["snapshots"]]
    assert notes == ["initial cap from charter", "Series C closing 2025-06-01"]
