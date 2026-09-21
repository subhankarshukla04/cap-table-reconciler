"""Snapshot-stamped live formula workbook (W3.5 / closes GAP-32)."""

from __future__ import annotations

import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.formula_workbook import build_formula_workbook
from src.models import (
    AntiDilution,
    AntiDilutionVariant,
    CapTable,
    Company,
    LiquidationPreference,
    LPType,
    ShareClass,
    ShareClassType,
)
from src.waterfall import compute_waterfall


def _ct():
    return CapTable(
        company=Company(name="Stamp Co", currency="USD"),
        share_classes=[
            ShareClass(name="C", type=ShareClassType.common, shares_outstanding=1000),
            ShareClass(
                name="A", type=ShareClassType.preferred, shares_outstanding=1000,
                issue_price=1.0, issue_date=date(2024, 1, 1), seniority_rank=1,
                liquidation_preference=LiquidationPreference(
                    multiple=1, amount=1000, type=LPType.non_participating),
                anti_dilution=AntiDilution(variant=AntiDilutionVariant.broad_based_weighted_average),
            ),
        ],
    )


def test_workbook_without_stamp_has_no_stamp_sheet():
    """Existing callers (no snapshot_stamp arg) must continue to work."""
    ct = _ct()
    wb = build_formula_workbook(ct, compute_waterfall(ct))
    assert "Snapshot Stamp" not in wb.sheetnames


def test_workbook_with_stamp_records_metadata():
    ct = _ct()
    stamp = {
        "engagement_id": "eng-uuid",
        "snapshot_id": "snap-uuid",
        "memo_version": "1.0",
        "engine_version": "abcd1234",
        "pack_version": "v2026.3.0",
        "generated_at": "2026-05-25T10:00:00+00:00",
    }
    wb = build_formula_workbook(ct, compute_waterfall(ct), snapshot_stamp=stamp)
    assert "Snapshot Stamp" in wb.sheetnames
    ws = wb["Snapshot Stamp"]
    # Field/Value pairs in column A/B.
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(1, 10)}
    assert values["engagement_id"] == "eng-uuid"
    assert values["snapshot_id"] == "snap-uuid"
    assert values["memo_version"] == "1.0"
    assert values["pack_version"] == "v2026.3.0"


def test_stamp_sheet_is_hidden():
    ct = _ct()
    wb = build_formula_workbook(
        ct, compute_waterfall(ct),
        snapshot_stamp={
            "engagement_id": "x", "snapshot_id": "y",
            "generated_at": "2026-05-25T10:00:00+00:00",
        },
    )
    ws = wb["Snapshot Stamp"]
    assert ws.sheet_state == "hidden"


def test_stamp_round_trips_through_xlsx():
    ct = _ct()
    stamp = {
        "engagement_id": "X1", "snapshot_id": "S1",
        "generated_at": "2026-05-25T10:00:00+00:00",
    }
    wb = build_formula_workbook(ct, compute_waterfall(ct), snapshot_stamp=stamp)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    reopened = openpyxl.load_workbook(buf)
    assert "Snapshot Stamp" in reopened.sheetnames
    ws = reopened["Snapshot Stamp"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(1, 10)}
    assert values["engagement_id"] == "X1"
    assert values["snapshot_id"] == "S1"
