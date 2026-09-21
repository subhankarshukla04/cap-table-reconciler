"""Volatility input pack tests (W3.4)."""

from __future__ import annotations

import tempfile
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.opm.backsolve import MarketInputs
from src.opm.vol_pack import build_vol_pack_template, read_vol_pack


def _fill(wb, slug_value_pairs: dict[str, object]):
    """Fill the template's value cells by looking up their slug in column D."""
    ws = wb["Market Inputs"]
    for r in range(1, 200):
        slug = ws.cell(row=r, column=4).value
        if slug and slug in slug_value_pairs:
            ws.cell(row=r, column=2, value=slug_value_pairs[slug])


def _save(wb) -> Path:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        wb.save(fh.name)
    return Path(fh.name)


def test_template_has_required_sheets():
    wb = build_vol_pack_template()
    assert "Market Inputs" in wb.sheetnames
    assert "Read me" in wb.sheetnames


def test_unfilled_required_fields_block_readback():
    wb = build_vol_pack_template()
    p = _save(wb)
    try:
        result = read_vol_pack(p)
        assert not result.is_valid
        # Required sourcing fields blank by default → in missing
        assert any("peer-set tickers" in s for s in result.missing_required)
    finally:
        p.unlink(missing_ok=True)


def test_fully_filled_template_produces_market_inputs():
    wb = build_vol_pack_template()
    _fill(wb, {
        "volatility": 0.60,
        "time_to_liquidity_years": 5.0,
        "risk_free_rate": 0.04,
        "dividend_yield": 0.0,
        "dlom": 0.30,
        "vol_peer_tickers": "AAPL, MSFT, GOOG",
        "vol_window_months": "24",
        "vol_size_adjustment_bps": "250",
        "vol_citation": "Damodaran 2025 — implied equity vol",
        "time_basis": "Board guidance",
        "rfr_source": "US Treasury 5y, 2026-05-25",
        "dlom_basis": "Finnerty 2012",
    })
    p = _save(wb)
    try:
        result = read_vol_pack(p)
        assert result.is_valid
        assert result.market == MarketInputs(
            volatility=0.60,
            time_to_liquidity_years=5.0,
            risk_free_rate=0.04,
            dividend_yield=0.0,
            dlom=0.30,
        )
        assert "AAPL" in result.sourcing["vol_peer_tickers"]
        assert "Damodaran" in result.sourcing["vol_citation"]
    finally:
        p.unlink(missing_ok=True)


def test_non_numeric_value_in_required_field_flagged():
    wb = build_vol_pack_template()
    _fill(wb, {
        "volatility": "not a number",
        "vol_peer_tickers": "AAPL",
        "vol_window_months": "24",
        "vol_citation": "any",
        "time_basis": "any",
        "rfr_source": "any",
        "dlom_basis": "any",
    })
    p = _save(wb)
    try:
        result = read_vol_pack(p)
        assert not result.is_valid
        assert any("not numeric" in s for s in result.missing_required)
    finally:
        p.unlink(missing_ok=True)


def test_bytesio_input_works():
    wb = build_vol_pack_template()
    _fill(wb, {
        "volatility": 0.5, "time_to_liquidity_years": 4.0, "risk_free_rate": 0.045,
        "vol_peer_tickers": "X", "vol_window_months": "12",
        "vol_citation": "x", "time_basis": "x", "rfr_source": "x", "dlom_basis": "x",
    })
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = read_vol_pack(buf)
    assert result.is_valid
