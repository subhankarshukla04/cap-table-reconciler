"""Volatility input pack template (W3.4 / SYSTEM_SPEC §5.1 + §7).

A structured Excel template the analyst fills with their volatility,
time-to-liquidity, risk-free-rate, dividend-yield, and DLOM choices,
PLUS the sourcing fields (peer ticker(s), observation window, size
adjustment, citation). The tool reads back the populated template and
hands a validated MarketInputs to OPM Backsolve.

Hard refusal: OPM Backsolve cannot run if any required sourcing field
is empty. The spec is explicit — vol/time/rfr/DLOM are JUDGMENT inputs;
the tool surfaces them, the analyst defends them. A blank sourcing
field is a fabrication-shaped failure mode (the analyst sets a number
without saying where it came from), and the tool refuses.

Spec promise (§7 out-of-scope refusals):
  - "No volatility peer-set engine. If built, it is a separate research
     project, not bolted into the reconciler."

This module honours that — it does NOT compute vol. It just structures
the form so the analyst's choices are auditable.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .backsolve import MarketInputs


# ---- The fillable template -------------------------------------------------


_FIELDS = [
    # (name, label, default_hint, required, group)
    ("volatility",            "Annualised volatility (decimal, e.g. 0.55)", 0.55,  True,  "Inputs"),
    ("time_to_liquidity_years","Time to liquidity (years)",                  4.0,   True,  "Inputs"),
    ("risk_free_rate",        "Continuously-compounded risk-free rate",     0.045, True,  "Inputs"),
    ("dividend_yield",        "Continuous dividend yield",                  0.0,   False, "Inputs"),
    ("dlom",                  "DLOM applied to common (decimal, e.g. 0.25)",0.25,  False, "Inputs"),
]

_SOURCING = [
    # (name, label, required)
    ("vol_peer_tickers",      "Volatility peer-set tickers (comma-separated)", True),
    ("vol_window_months",     "Observation window for vol (months)",          True),
    ("vol_size_adjustment_bps","Size adjustment (basis points)",              False),
    ("vol_citation",          "Citation / source (e.g. Damodaran 2025, Bloomberg pull date)", True),
    ("time_basis",            "Basis for time-to-liquidity (e.g. PWERM, board guidance)",     True),
    ("rfr_source",            "Risk-free-rate source (e.g. US Treasury yield curve, date)",   True),
    ("dlom_basis",            "DLOM basis (e.g. Finnerty option, restricted-stock studies)",  True),
]


_INPUT_FILL = PatternFill("solid", fgColor="FFF2C8")
_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
_HEADER_FONT = Font(bold=True)


@dataclass
class VolPackReadback:
    market: Optional[MarketInputs]
    sourcing: dict[str, str]
    missing_required: list[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return self.market is not None and not self.missing_required


# ---- Build the empty template ---------------------------------------------


def build_vol_pack_template() -> Workbook:
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    inputs = wb.create_sheet("Market Inputs")
    inputs["A1"] = "Volatility Input Pack"
    inputs["A1"].font = Font(bold=True, size=14)
    inputs["A2"] = (
        "Fill the yellow cells. The OPM Backsolve sibling will refuse to "
        "run if any REQUIRED sourcing field below is empty."
    )
    inputs["A2"].alignment = Alignment(wrap_text=True)
    inputs.row_dimensions[2].height = 36

    # Numeric inputs
    inputs["A4"] = "Field"
    inputs["B4"] = "Value"
    inputs["C4"] = "Required"
    for cell in ("A4", "B4", "C4"):
        inputs[cell].font = _HEADER_FONT
        inputs[cell].fill = _HEADER_FILL
    for i, (name, label, default, required, _group) in enumerate(_FIELDS, start=5):
        inputs.cell(row=i, column=1, value=label)
        v_cell = inputs.cell(row=i, column=2, value=default)
        v_cell.fill = _INPUT_FILL
        inputs.cell(row=i, column=3, value="Yes" if required else "No")
        inputs.cell(row=i, column=4, value=name)  # named-range slug

    # Sourcing fields
    base_row = 5 + len(_FIELDS) + 2
    inputs.cell(row=base_row, column=1, value="Sourcing field").font = _HEADER_FONT
    inputs.cell(row=base_row, column=2, value="Value").font = _HEADER_FONT
    inputs.cell(row=base_row, column=3, value="Required").font = _HEADER_FONT
    inputs.cell(row=base_row, column=1).fill = _HEADER_FILL
    inputs.cell(row=base_row, column=2).fill = _HEADER_FILL
    inputs.cell(row=base_row, column=3).fill = _HEADER_FILL
    for i, (name, label, required) in enumerate(_SOURCING, start=base_row + 1):
        inputs.cell(row=i, column=1, value=label)
        v_cell = inputs.cell(row=i, column=2, value="")
        v_cell.fill = _INPUT_FILL
        inputs.cell(row=i, column=3, value="Yes" if required else "No")
        inputs.cell(row=i, column=4, value=name)

    for col_idx, width in enumerate([50, 20, 12, 30], start=1):
        inputs.column_dimensions[get_column_letter(col_idx)].width = width

    # Read-me
    rm = wb.create_sheet("Read me")
    rm["A1"] = "How to fill this pack"
    rm["A1"].font = Font(bold=True, size=14)
    rm["A3"] = (
        "1. Fill every yellow cell. Required fields must be non-empty.\n"
        "2. Citation fields must point to a defensible source (Damodaran, "
        "Bloomberg, AICPA practice aid, peer-comparable bracket). The tool "
        "stores your text verbatim; it does NOT validate the source.\n"
        "3. Save and pass back to the engine. The OPM Backsolve sibling "
        "calls read_vol_pack() on this file and refuses to run if any "
        "required field is empty."
    )
    rm["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    rm.row_dimensions[3].height = 110
    rm.column_dimensions["A"].width = 80

    return wb


def build_vol_pack_template_bytes() -> bytes:
    wb = build_vol_pack_template()
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Read a filled pack back ----------------------------------------------


def _cell_value(ws, name: str, max_row: int = 200) -> Optional[object]:
    """Look up a field by its slug stored in column D."""
    for r in range(1, max_row + 1):
        slug = ws.cell(row=r, column=4).value
        if slug == name:
            return ws.cell(row=r, column=2).value
    return None


def read_vol_pack(path: Path | str | BytesIO) -> VolPackReadback:
    wb = load_workbook(path, data_only=True)
    if "Market Inputs" not in wb.sheetnames:
        return VolPackReadback(
            market=None, sourcing={}, missing_required=["Market Inputs sheet"]
        )
    ws = wb["Market Inputs"]

    # W3-AUDIT B1: numeric bounds. SYSTEM_SPEC §5.1 expects vol in [0, 5],
    # time > 0, rf in [-0.05, 1], dlom in [0, 1], dividend_yield in [0, 1].
    # A negative volatility or > 100% DLOM is a transcription error; refuse.
    _BOUNDS = {
        "volatility": (0.0, 5.0),
        "time_to_liquidity_years": (0.0, 50.0),  # > 0 enforced below
        "risk_free_rate": (-0.05, 1.0),
        "dividend_yield": (0.0, 1.0),
        "dlom": (0.0, 1.0),
    }
    numeric_values: dict[str, float] = {}
    missing: list[str] = []
    for name, label, _default, required, _group in _FIELDS:
        raw = _cell_value(ws, name)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if required:
                missing.append(label)
            continue
        try:
            v = float(raw)
        except (TypeError, ValueError):
            missing.append(f"{label} (not numeric: {raw!r})")
            continue
        lo, hi = _BOUNDS[name]
        if not (lo <= v <= hi):
            missing.append(
                f"{label} (value {v} outside expected range [{lo}, {hi}])"
            )
            continue
        if name == "time_to_liquidity_years" and v <= 0:
            missing.append(f"{label} (must be > 0)")
            continue
        numeric_values[name] = v

    sourcing_values: dict[str, str] = {}
    for name, label, required in _SOURCING:
        raw = _cell_value(ws, name)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if required:
                missing.append(label)
            sourcing_values[name] = ""
        else:
            sourcing_values[name] = str(raw).strip()

    market = None
    if not missing and {"volatility", "time_to_liquidity_years", "risk_free_rate"} <= set(numeric_values):
        market = MarketInputs(
            volatility=numeric_values["volatility"],
            time_to_liquidity_years=numeric_values["time_to_liquidity_years"],
            risk_free_rate=numeric_values["risk_free_rate"],
            dividend_yield=numeric_values.get("dividend_yield", 0.0),
            dlom=numeric_values.get("dlom", 0.0),
        )

    return VolPackReadback(
        market=market,
        sourcing=sourcing_values,
        missing_required=missing,
    )
